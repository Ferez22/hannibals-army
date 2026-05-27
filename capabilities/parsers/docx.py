"""DOCX parser via python-docx. Extracts paragraphs + tables."""
from __future__ import annotations

from pathlib import Path


def parse(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts: list[str] = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
