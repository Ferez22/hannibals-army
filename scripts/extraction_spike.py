"""
Extraction spike — Phase 0.5

Measures Gemma4:e2b extraction quality on real QartMina docs.
Throwaway script. Decides whether plan needs adjustment before Phase 0.

Run:
    .venv/bin/python scripts/extraction_spike.py

Reads:
    data/samples/*.{pdf,docx,xlsx}
    data/samples/ground_truth.yml

Writes:
    docs/EXTRACTION_BASELINE.md (report)
"""
from __future__ import annotations

import json
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from langchain_ollama.llms import OllamaLLM

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
REPO = Path(__file__).resolve().parent.parent
SAMPLES_DIR = REPO / "data" / "samples"
GROUND_TRUTH = SAMPLES_DIR / "ground_truth.yml"
REPORT_PATH = REPO / "docs" / "EXTRACTION_BASELINE.md"
MODEL = "gemma4:e2b"
MAX_CHARS = 8000  # cap per doc for spike — keep within gemma4:e2b context

EXTRACTION_PROMPT = """You extract entities for a company knowledge graph. Return ONLY valid JSON, no prose, no markdown fences.

ENTITY DEFINITIONS (strict — when in doubt, omit):

- person: a named human being. Skip generic references ("the team", "the speaker", "the client").
- team: a named subdivision INSIDE a company (e.g., "Engineering", "HR", "Backend"). Do NOT extract the company itself as a team. Do NOT extract job titles. A team must be an internal organizational unit.
- project: a named initiative or product the company works on (e.g., "Additionality", "Product Alpha"). Do NOT extract todo items, tasks, or to-do list entries as projects. A project has a name that would appear on a roadmap, not on a checklist.
- rule: an explicit policy, process, or value, typically marked by words like "policy", "must", "shall", "rule", "process". Skip incidental statements. If the doc isn't explicitly defining a rule, omit.
- event: a company-level dated occurrence (meeting, launch, signing, all-hands). Do NOT extract personal life events (births, passport issuance, ID card issuance). A date alone is not an event.
- edge: relationship between two extracted entities. Valid types: MEMBER_OF, RUNS, PARTICIPATED_IN, AUTHORED, OWNED_BY, CHILD_OF.

OUTPUT RULES:
- Only extract what is EXPLICITLY in the document text
- Use exact names as they appear in the document
- If no entity of a type exists, return an empty list []
- Do not infer or invent

OUTPUT SCHEMA:
{
  "persons":  [{"name": "...", "email": "..."|null, "role": "..."|null}],
  "teams":    [{"name": "...", "parent": null}],
  "projects": [{"name": "...", "lead": "..."|null, "status": "..."|null}],
  "rules":    [{"title": "...", "category": "policy|process|value"}],
  "events":   [{"name": "...", "date": "YYYY-MM-DD"|null}],
  "edges":    [{"from": "...", "type": "MEMBER_OF", "to": "...", "role": "..."|null}]
}

DOCUMENT:
---
__DOC_TEXT__
---

JSON:"""


# ----------------------------------------------------------------------------
# Parsers
# ----------------------------------------------------------------------------
def parse_pdf(path: Path) -> str:
    import pymupdf

    parts: list[str] = []
    with pymupdf.open(path) as doc:
        for page in doc:
            parts.append(page.get_text())
    return "\n".join(parts)


def parse_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def parse_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


PARSERS = {".pdf": parse_pdf, ".docx": parse_docx, ".xlsx": parse_xlsx}


def parse_doc(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in PARSERS:
        raise ValueError(f"Unsupported format: {ext}")
    text = PARSERS[ext](path)
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n[TRUNCATED]"
    return text


# ----------------------------------------------------------------------------
# Normalization (for matching extracted vs truth)
# ----------------------------------------------------------------------------
def normalize_name(s: str | None) -> frozenset[str]:
    """Return a frozenset of name tokens — order-insensitive matching.

    "Ali Khribi" → {"ali", "khribi"}
    "KHRIBI ALI" → {"ali", "khribi"}  (matches)
    "Smith, John" → {"john", "smith"} (matches "John Smith")
    """
    if not s:
        return frozenset()
    s = unicodedata.normalize("NFKD", s)
    s = s.encode("ascii", "ignore").decode("ascii")
    s = s.lower()
    s = re.sub(r"\b(dr|mr|mrs|ms|prof|jr|sr|phd|md)\.?\b", "", s)
    s = re.sub(r"[^\w\s]", " ", s)
    tokens = {t for t in s.split() if len(t) > 1}  # drop single-letter noise
    return frozenset(tokens)


# ----------------------------------------------------------------------------
# Extraction call
# ----------------------------------------------------------------------------
@dataclass
class ExtractionResult:
    raw_text: str
    parsed: dict[str, Any] | None
    json_valid: bool
    error: str | None
    latency_s: float


def extract(llm: OllamaLLM, text: str) -> ExtractionResult:
    prompt = EXTRACTION_PROMPT.replace("__DOC_TEXT__", text)
    start = time.time()
    try:
        raw = llm.invoke(prompt)
    except Exception as e:
        return ExtractionResult("", None, False, f"LLM error: {e}", time.time() - start)
    latency = time.time() - start

    # strip markdown fences if model adds them
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)

    try:
        parsed = json.loads(cleaned)
        return ExtractionResult(raw, parsed, True, None, latency)
    except json.JSONDecodeError as e:
        # try to find a JSON object inside
        match = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if match:
            try:
                parsed = json.loads(match.group(0))
                return ExtractionResult(raw, parsed, True, f"recovered (was: {e})", latency)
            except json.JSONDecodeError as e2:
                return ExtractionResult(raw, None, False, f"JSON invalid: {e2}", latency)
        return ExtractionResult(raw, None, False, f"JSON invalid: {e}", latency)


# ----------------------------------------------------------------------------
# Scoring
# ----------------------------------------------------------------------------
@dataclass
class TypeScore:
    truth_count: int = 0
    extracted_count: int = 0
    matched: int = 0

    @property
    def precision(self) -> float:
        return self.matched / self.extracted_count if self.extracted_count else 0.0

    @property
    def recall(self) -> float:
        return self.matched / self.truth_count if self.truth_count else 0.0


def score_by_name(truth: list[dict], extracted: list[dict], key: str = "name") -> TypeScore:
    """Match by token-set: truth and extracted match if their token sets overlap by
    at least the smaller set's size (i.e., one is contained in the other, or equal)."""
    truth_sets = [normalize_name(t.get(key, "")) for t in truth if t.get(key)]
    extracted_sets = [normalize_name(e.get(key, "")) for e in extracted if e.get(key)]
    truth_sets = [s for s in truth_sets if s]
    extracted_sets = [s for s in extracted_sets if s]

    matched_truth: set[int] = set()
    matched_extracted: set[int] = set()
    for i, ts in enumerate(truth_sets):
        for j, es in enumerate(extracted_sets):
            if j in matched_extracted:
                continue
            overlap = ts & es
            min_size = min(len(ts), len(es))
            if min_size > 0 and len(overlap) >= min_size:
                matched_truth.add(i)
                matched_extracted.add(j)
                break

    return TypeScore(len(truth_sets), len(extracted_sets), len(matched_truth))


def score_rules(truth: list[dict], extracted: list[dict]) -> TypeScore:
    return score_by_name(truth, extracted, key="title")


@dataclass
class DocReport:
    filename: str
    json_valid: bool
    error: str | None
    latency_s: float
    scores: dict[str, TypeScore] = field(default_factory=dict)


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main() -> int:
    if not GROUND_TRUTH.exists():
        print(f"ERROR: {GROUND_TRUTH} not found")
        return 1

    truth_data = yaml.safe_load(GROUND_TRUTH.read_text())
    llm = OllamaLLM(model=MODEL, temperature=0.0)

    print(f"Model: {MODEL}")
    print(f"Samples dir: {SAMPLES_DIR}")
    print(f"Docs to process: {len(truth_data)}\n")

    reports: list[DocReport] = []

    for filename, truth in truth_data.items():
        path = SAMPLES_DIR / filename
        print(f"[{filename}]")

        if not path.exists():
            print(f"  SKIP — file not found")
            reports.append(DocReport(filename, False, "file not found", 0.0))
            continue

        ext = path.suffix.lower()
        if ext not in PARSERS:
            print(f"  SKIP — unsupported format {ext}")
            reports.append(DocReport(filename, False, f"unsupported {ext}", 0.0))
            continue

        try:
            text = parse_doc(path)
        except Exception as e:
            print(f"  PARSE FAILED: {e}")
            reports.append(DocReport(filename, False, f"parse error: {e}", 0.0))
            continue

        if not text.strip():
            print(f"  SKIP — empty text after parsing (likely scanned/image PDF)")
            reports.append(DocReport(filename, False, "empty text extracted", 0.0))
            continue

        print(f"  text: {len(text)} chars")
        result = extract(llm, text)
        print(f"  latency: {result.latency_s:.1f}s, json_valid: {result.json_valid}")

        if not result.json_valid or result.parsed is None:
            print(f"  ERROR: {result.error}")
            reports.append(DocReport(filename, False, result.error, result.latency_s))
            continue

        scores: dict[str, TypeScore] = {}
        e = result.parsed
        scores["persons"] = score_by_name(truth.get("expected_persons", []), e.get("persons", []))
        scores["teams"] = score_by_name(truth.get("expected_teams", []), e.get("teams", []))
        scores["projects"] = score_by_name(truth.get("expected_projects", []), e.get("projects", []))
        scores["rules"] = score_rules(truth.get("expected_rules", []), e.get("rules", []))
        scores["events"] = score_by_name(truth.get("expected_events", []), e.get("events", []))

        for t, s in scores.items():
            if s.truth_count or s.extracted_count:
                print(f"  {t:10s}: P={s.precision:.2f} R={s.recall:.2f}  (truth={s.truth_count}, extracted={s.extracted_count}, matched={s.matched})")

        reports.append(DocReport(filename, True, result.error, result.latency_s, scores))
        # Dump raw extraction for inspection
        dump_path = REPO / "data" / "samples" / f"_extracted_{path.stem}.json"
        dump_path.write_text(json.dumps(result.parsed, indent=2, ensure_ascii=False))
        print(f"  raw → {dump_path.name}")
        print()

    write_report(reports)
    print(f"\nReport: {REPORT_PATH}")
    return 0


def write_report(reports: list[DocReport]) -> None:
    lines: list[str] = []
    lines.append("# Extraction Baseline — Phase 0.5 Results")
    lines.append("")
    lines.append(f"Model: `{MODEL}`")
    lines.append(f"Docs processed: {len(reports)}")
    lines.append("")

    # Aggregate
    json_valid_count = sum(1 for r in reports if r.json_valid)
    json_validity = json_valid_count / len(reports) if reports else 0.0
    lines.append(f"## JSON Validity: {json_validity:.0%} ({json_valid_count}/{len(reports)})")
    lines.append("")

    # Per-type aggregates (precision/recall across docs)
    type_agg: dict[str, TypeScore] = {}
    for r in reports:
        for t, s in r.scores.items():
            agg = type_agg.setdefault(t, TypeScore())
            agg.truth_count += s.truth_count
            agg.extracted_count += s.extracted_count
            agg.matched += s.matched

    lines.append("## Aggregate Scores")
    lines.append("")
    lines.append("| Type | Precision | Recall | Truth | Extracted | Matched |")
    lines.append("|------|----------:|-------:|------:|----------:|--------:|")
    for t, s in type_agg.items():
        lines.append(f"| {t} | {s.precision:.2f} | {s.recall:.2f} | {s.truth_count} | {s.extracted_count} | {s.matched} |")
    lines.append("")

    # Decision gates
    person_score = type_agg.get("persons", TypeScore())
    team_score = type_agg.get("teams", TypeScore())
    lines.append("## Decision Gates")
    lines.append("")
    lines.append(f"- JSON validity ≥ 80%: **{'PASS' if json_validity >= 0.8 else 'FAIL'}** ({json_validity:.0%})")
    lines.append(f"- Person+Team precision ≥ 70%: persons={person_score.precision:.0%} teams={team_score.precision:.0%}")
    lines.append(f"- Person+Team recall ≥ 60%: persons={person_score.recall:.0%} teams={team_score.recall:.0%}")
    lines.append("")

    lines.append("## Per-Document Results")
    lines.append("")
    for r in reports:
        lines.append(f"### `{r.filename}`")
        lines.append(f"- JSON valid: {r.json_valid}")
        lines.append(f"- Latency: {r.latency_s:.1f}s")
        if r.error:
            lines.append(f"- Note: {r.error}")
        if r.scores:
            for t, s in r.scores.items():
                if s.truth_count or s.extracted_count:
                    lines.append(f"  - {t}: P={s.precision:.2f} R={s.recall:.2f} (truth={s.truth_count}, extracted={s.extracted_count})")
        lines.append("")

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text("\n".join(lines))


if __name__ == "__main__":
    sys.exit(main())
